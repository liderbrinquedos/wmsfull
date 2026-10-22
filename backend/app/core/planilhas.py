"""Leitura das planilhas de referencia (fichas tecnicas e localizacao).

Formato real medido em 2026-10-22:

- `FICHA TECNICA *.xlsx`: cabecalho na linha 2, juncao com o catalogo por `REF.`.
  Variacoes: `DUN14` (APOLO), `DUN INT` (LIDER), `DESCRICAO` ausente (FANDOMBOX).
- `Localizacao.xlsx`: cabecalho na linha 1, juncao por `Codigo`.

O parser nao conhece banco: devolve linhas validadas e a lista de problemas de
cada uma. Quem grava e o router de importacao.
"""
from __future__ import annotations

import unicodedata
from dataclasses import dataclass, field
from io import BytesIO
from typing import Iterable

from openpyxl import load_workbook

from app.core.barcodes import ean13_valido, ean14_valido

# Nomes reconhecidos por campo, em ordem de preferencia.
COLUNAS = {
    "ref": ("REF",),
    "descricao": ("DESCRICAO",),
    "cx_master": ("CX MASTER",),
    "ean13": ("EAN 13", "EAN13", "EAN"),
    "dun14": ("DUN 14", "DUN14", "DUN INT"),
    "codigo": ("CODIGO",),
    "qtd_embalagem": ("QUANTIDADE DE EMBALAGENS", "QTD EMBALAGEM", "QTD EMBALAGENS"),
    "localizacao": ("LOCALIZACAO",),
}

# Colunas que a planilha precisa ter (as demais sao opcionais).
OBRIGATORIAS = {
    "ficha": {"ref"},
    "localizacao": {"codigo"},
}


@dataclass
class LinhaFicha:
    linha: int
    ref: str
    descricao: str | None = None
    ean13: str | None = None
    dun14: str | None = None
    cx_master: int | None = None
    problemas: list[str] = field(default_factory=list)


@dataclass
class LinhaLocalizacao:
    linha: int
    codigo: str
    localizacao: str | None = None
    qtd_embalagem: int | None = None
    referencia: str | None = None
    problemas: list[str] = field(default_factory=list)


class PlanilhaInvalida(Exception):
    pass


def _normalizar(valor) -> str:
    if valor is None:
        return ""
    texto = str(valor).strip().upper()
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(caractere for caractere in texto if not unicodedata.combining(caractere))
    texto = texto.replace(".", " ")
    return " ".join(texto.split())


def _texto(valor) -> str | None:
    if valor is None:
        return None
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    texto = str(valor).strip()
    return texto or None


def _codigo_barras(valor) -> str | None:
    """EAN pode vir como numero na planilha; evita '7.89123e+12'."""
    if valor is None:
        return None
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    texto = str(valor).strip()
    if not texto:
        return None
    if texto.endswith(".0"):
        texto = texto[:-2]
    return texto


def _inteiro(valor) -> int | None:
    if valor is None or valor == "":
        return None
    if isinstance(valor, (int, float)):
        return int(valor)
    try:
        return int(float(str(valor).replace(",", ".")))
    except ValueError:
        return None


def _abrir(conteudo: bytes):
    try:
        return load_workbook(BytesIO(conteudo), read_only=True, data_only=True)
    except Exception as exc:
        raise PlanilhaInvalida(f"nao foi possivel abrir o arquivo como .xlsx: {exc}") from exc


def _linhas(planilha) -> list[list]:
    aba = planilha[planilha.sheetnames[0]]
    return [list(linha) for linha in aba.iter_rows(values_only=True)]


def _indices_cabecalho(linhas: list[list], tipo: str) -> tuple[int, dict[str, int]]:
    """Acha a linha de cabecalho e mapeia campo -> indice da coluna."""
    exigidas = OBRIGATORIAS[tipo]
    for indice, linha in enumerate(linhas[:10]):
        rotulos = [_normalizar(celula) for celula in linha]
        mapa: dict[str, int] = {}
        for campo, aliases in COLUNAS.items():
            for alias in aliases:
                if alias in rotulos:
                    mapa[campo] = rotulos.index(alias)
                    break
        if exigidas <= mapa.keys():
            return indice, mapa
    raise PlanilhaInvalida(
        "cabecalho nao encontrado: a planilha precisa de "
        + ", ".join(sorted(exigidas))
        + " nas primeiras 10 linhas"
    )


def _celulas(linha: list, mapa: dict[str, int]) -> dict[str, object]:
    def pegar(campo: str):
        indice = mapa.get(campo)
        if indice is None or indice >= len(linha):
            return None
        return linha[indice]

    return {campo: pegar(campo) for campo in mapa}


def ler_ficha(conteudo: bytes) -> list[LinhaFicha]:
    planilha = _abrir(conteudo)
    try:
        linhas = _linhas(planilha)
    finally:
        planilha.close()

    indice_cabecalho, mapa = _indices_cabecalho(linhas, "ficha")
    resultado: list[LinhaFicha] = []

    for numero, linha in enumerate(linhas[indice_cabecalho + 1 :], start=indice_cabecalho + 2):
        if not any(celula not in (None, "") for celula in linha):
            continue

        valores = _celulas(linha, mapa)
        ref = _texto(valores.get("ref"))
        if not ref:
            continue

        item = LinhaFicha(
            linha=numero,
            ref=ref,
            descricao=_texto(valores.get("descricao")),
            ean13=_codigo_barras(valores.get("ean13")),
            dun14=_codigo_barras(valores.get("dun14")),
            cx_master=_inteiro(valores.get("cx_master")),
        )

        if item.ean13 and not ean13_valido(item.ean13):
            item.problemas.append("EAN13 com digito verificador invalido")
            item.ean13 = None
        if item.dun14 and not ean14_valido(item.dun14):
            item.problemas.append("DUN14 com digito verificador ou tamanho invalido")
            item.dun14 = None

        resultado.append(item)

    return resultado


def ler_localizacao(conteudo: bytes) -> list[LinhaLocalizacao]:
    planilha = _abrir(conteudo)
    try:
        linhas = _linhas(planilha)
    finally:
        planilha.close()

    indice_cabecalho, mapa = _indices_cabecalho(linhas, "localizacao")
    resultado: list[LinhaLocalizacao] = []

    for numero, linha in enumerate(linhas[indice_cabecalho + 1 :], start=indice_cabecalho + 2):
        if not any(celula not in (None, "") for celula in linha):
            continue

        valores = _celulas(linha, mapa)
        codigo = _texto(valores.get("codigo"))
        if not codigo:
            continue

        item = LinhaLocalizacao(
            linha=numero,
            codigo=codigo,
            localizacao=_texto(valores.get("localizacao")),
            qtd_embalagem=_inteiro(valores.get("qtd_embalagem")),
            referencia=_texto(valores.get("referencia")),
        )
        if item.localizacao and len(item.localizacao) > 40:
            item.problemas.append("localizacao com mais de 40 caracteres")
            item.localizacao = None

        resultado.append(item)

    return resultado


def contar_duplicados(valores: Iterable[str]) -> dict[str, int]:
    contagem: dict[str, int] = {}
    for valor in valores:
        contagem[valor] = contagem.get(valor, 0) + 1
    return {valor: total for valor, total in contagem.items() if total > 1}
